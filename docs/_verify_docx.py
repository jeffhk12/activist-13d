"""Sanity-check the generated design record. Tier: 0 (scratch)."""
import os
from docx import Document

P = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                 "Activist-13D-Design-Record.docx")
d = Document(P)
print("paragraphs:", len(d.paragraphs), " tables:", len(d.tables))
print("size KB:", round(os.path.getsize(P) / 1024, 1))
print("--- outline ---")
for p in d.paragraphs:
    if p.style.name.startswith("Heading"):
        lvl = p.style.name.split()[-1]
        print(("  " * (int(lvl) - 1)) + f"H{lvl}  {p.text}")
