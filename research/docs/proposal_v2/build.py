"""Render the canonical V2 Markdown proposal into an editable two-page DOCX.

Inputs: docs/PROFESSOR-PROPOSAL-v2.md. Output: matching DOCX in docs/.
Identity: paragraphs preserve source text; no estimates or data transforms.
Tier: 0
Run: python docs/proposal_v2/build.py (requires python-docx 1.2.0).
"""

import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

DOCS = Path(__file__).resolve().parents[1]
source = DOCS / "PROFESSOR-PROPOSAL-v2.md"
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.65)
section.left_margin = section.right_margin = Inches(0.8)
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.04
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.widow_control = True
for name, size in [("Title", 18), ("Subtitle", 11), ("Heading 1", 12)]:
    style = doc.styles[name]
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = name != "Subtitle"
    style.paragraph_format.space_before = Pt(8 if name == "Heading 1" else 0)
    style.paragraph_format.space_after = Pt(5)

for block in source.read_text().strip().split("\n\n"):
    if block == "<!-- PAGEBREAK -->":
        doc.add_page_break()
        continue
    if block.startswith("# "):
        doc.add_paragraph(block[2:], "Title")
        continue
    if block.startswith("## "):
        doc.add_paragraph(block[3:], "Heading 1")
        continue
    p = doc.add_paragraph()
    if block.startswith(("Project Proposal", "Jeff Lu |", "For Professor")):
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
    for i, text in enumerate(re.split(r"\*\*(.*?)\*\*", block)):
        r = p.add_run(text)
        r.bold = bool(i % 2)
    if block.startswith("For Professor"):
        p.paragraph_format.space_after = Pt(8)

footer = section.footer.paragraphs[0]
footer.alignment = 2
footer.paragraph_format.space_after = Pt(0)
run = footer.add_run()
run.font.name = "Times New Roman"
run.font.size = Pt(9)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
run._r.addnext(field)
doc.core_properties.author = "Jeff Lu"
doc.core_properties.title = "Activist Intent and Investment Selection"
doc.core_properties.subject = "Project proposal for Professor Dmitry Udler"
doc.core_properties.comments = "Developed and drafted with ChatGPT and Codex assistance."
# Remove inherited template rules and theme fonts so the intended typography
# and border-free title survive conversion in both Word and LibreOffice.
for style in doc.styles:
    for border in list(style.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    if style.name in {"Normal", "Title", "Subtitle", "Heading 1"}:
        fonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
        if fonts is not None:
            for attribute in list(fonts.attrib):
                if "Theme" in attribute:
                    del fonts.attrib[attribute]
            for attribute in ("ascii", "hAnsi", "cs", "eastAsia"):
                fonts.set(qn("w:" + attribute), "Times New Roman")
        for color in style.element.iter(qn("w:color")):
            for attribute in list(color.attrib):
                if "theme" in attribute.lower():
                    del color.attrib[attribute]
            color.set(qn("w:val"), "000000")
output = source.with_suffix(".docx")
doc.save(output)
print(output)
